"""Build docs/presentations/DCU-Active-IT-Danamon-Risk-Briefing.pptx (briefing for Bank Danamon IT).

Content mirrors docs/presentations/DCU-Active-IT-Danamon-Risk-Briefing.md. Requires python-pptx.
Usage: python scripts/build-it-briefing-pptx.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

RED = RGBColor(0xC8, 0x10, 0x2E)
RED_DARK = RGBColor(0x7F, 0x1D, 0x1D)
RED_SOFT = RGBColor(0xFD, 0xEC, 0xEE)
INK = RGBColor(0x1F, 0x29, 0x37)
GREY = RGBColor(0x6B, 0x72, 0x80)
GREY_LIGHT = RGBColor(0xF3, 0xF4, 0xF6)
LINE = RGBColor(0xD1, 0xD5, 0xDB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
AMBER = RGBColor(0xB4, 0x53, 0x09)
GREEN = RGBColor(0x15, 0x80, 0x3D)
FONT = 'Segoe UI'

OUT = Path(__file__).resolve().parent.parent / 'docs' / 'presentations' / 'DCU-Active-IT-Danamon-Risk-Briefing.pptx'

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
W, H = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]
TOTAL = 16


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    return s


def text(slide, x, y, w, h, value, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    lines = value if isinstance(value, list) else [value]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run(p, line, size, color, bold)
    return tb


def run(p, value, size, color, bold=False):
    """Adds text to a paragraph; **bold** segments are rendered bold."""
    parts = value.split('**')
    for i, part in enumerate(parts):
        if not part:
            continue
        r = p.add_run()
        r.text = part
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.bold = bold or (i % 2 == 1)


def bullets(slide, x, y, w, h, items, size=17, color=INK, gap=8, numbered=False):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        sub = item.startswith('  ')
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        pPr = p._p.get_or_add_pPr()
        indent = Emu(Inches(0.28 if not sub else 0.25))
        pPr.set('marL', str(Inches(0.3 if not sub else 0.65)))
        pPr.set('indent', str(-indent))
        bu_clr = pPr.makeelement(qn('a:buClr'), {})
        srgb = bu_clr.makeelement(qn('a:srgbClr'), {'val': 'C8102E' if not sub else '6B7280'})
        bu_clr.append(srgb)
        pPr.append(bu_clr)
        if numbered and not sub:
            pPr.append(pPr.makeelement(qn('a:buFont'), {'typeface': FONT}))
            pPr.append(pPr.makeelement(qn('a:buAutoNum'), {'type': 'arabicPeriod'}))
        else:
            pPr.append(pPr.makeelement(qn('a:buFont'), {'typeface': 'Arial'}))
            pPr.append(pPr.makeelement(qn('a:buChar'), {'char': '•' if not sub else '–'}))
        run(p, item.strip(), size - (2 if sub else 0), color if not sub else GREY)
    return tb


def base(title, n, kicker=None):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = WHITE
    rect(s, 0, 0, W, Inches(0.12), RED)
    if kicker:
        text(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.35), kicker.upper(), 12, RED, True)
    text(s, Inches(0.6), Inches(0.62), Inches(12.1), Inches(0.8), title, 30, INK, True)
    rect(s, Inches(0.6), Inches(1.42), Inches(0.9), Inches(0.06), RED)
    rect(s, 0, H - Inches(0.45), W, Inches(0.45), GREY_LIGHT)
    text(s, Inches(0.6), H - Inches(0.4), Inches(9), Inches(0.35),
         'DCU Active  ·  Arsitektur, Keamanan & Risiko  ·  Oktober 2026', 10, GREY, anchor=MSO_ANCHOR.MIDDLE)
    text(s, W - Inches(1.6), H - Inches(0.4), Inches(1.0), Inches(0.35), f'{n} / {TOTAL}', 10, GREY,
         align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    return s


def card(s, x, y, w, h, title, body, num=None, fill=GREY_LIGHT, title_color=INK, body_size=13):
    rect(s, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    rect(s, x, y + Inches(0.18), Inches(0.07), h - Inches(0.36), RED)
    tx = x + Inches(0.25)
    if num is not None:
        c = rect(s, tx, y + Inches(0.2), Inches(0.45), Inches(0.45), RED, shape=MSO_SHAPE.OVAL)
        tf = c.text_frame
        tf.margin_left = tf.margin_right = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run(p, str(num), 14, WHITE, True)
        text(s, tx + Inches(0.55), y + Inches(0.18), w - Inches(0.9), Inches(0.5), title, 15, title_color, True,
             anchor=MSO_ANCHOR.MIDDLE)
    else:
        text(s, tx, y + Inches(0.15), w - Inches(0.4), Inches(0.5), title, 15, title_color, True)
    text(s, tx, y + Inches(0.72), w - Inches(0.4), h - Inches(0.8), body, body_size, GREY)


# 1 — Title ------------------------------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
s.background.fill.solid()
s.background.fill.fore_color.rgb = RED
rect(s, 0, H - Inches(1.6), W, Inches(1.6), INK)
rect(s, Inches(0.8), Inches(1.7), Inches(1.2), Inches(0.08), WHITE)
text(s, Inches(0.8), Inches(1.95), Inches(11.5), Inches(1.2), 'DCU Active', 54, WHITE, True)
text(s, Inches(0.8), Inches(3.05), Inches(11.5), Inches(1.2),
     'Arsitektur Sistem, Keamanan, Risiko & Mitigasi', 30, WHITE)
text(s, Inches(0.8), Inches(4.05), Inches(11.5), Inches(0.6),
     'Platform booking fasilitas olahraga & rekreasi DCU (web app / PWA)', 18, RED_SOFT)
text(s, Inches(0.8), H - Inches(1.3), Inches(8), Inches(1.0),
     ['Presentasi kepada IT Bank Danamon', 'Oktober 2026'], 16, WHITE)
text(s, Inches(7.5), H - Inches(1.3), Inches(5.0), Inches(1.0),
     ['Product Owner: Ronny Iswara (DCU)', 'DCU@danamon.co.id'], 16, WHITE, align=PP_ALIGN.RIGHT)

# 2 — Objectives -------------------------------------------------------------------------------------------
s = base('Tujuan Sesi', 2, 'Pengantar')
goals = [
    ('Menjelaskan', 'Bagaimana DCU Active dibangun, di-hosting, dan diatur aksesnya.'),
    ('Memaparkan', 'Risiko teknis, keamanan, dan privasi beserta mitigasi yang sudah diterapkan.'),
    ('Membuktikan', 'Hasil pengujian sejauh ini. Staging aktif sejak 14 Oktober 2026.'),
    ('Meminta', 'Dukungan dan persetujuan IT Danamon untuk beberapa hal sebelum go-live.'),
]
for i, (t, b) in enumerate(goals):
    card(s, Inches(0.6 + i * 3.08), Inches(2.0), Inches(2.9), Inches(2.6), t, b, num=i + 1, body_size=15)
text(s, Inches(0.6), Inches(5.0), Inches(12), Inches(0.8),
     'Fokus: memberi IT Danamon gambaran yang cukup untuk menilai risiko dan menyetujui langkah menuju produksi.',
     15, GREY)

# 3 — What is it -------------------------------------------------------------------------------------------
s = base('Apa itu DCU Active?', 3, 'Ringkasan sistem')
bullets(s, Inches(0.6), Inches(1.8), Inches(7.2), Inches(5), [
    'Web app / PWA untuk **booking fasilitas DCU** dari HP atau laptop',
    'Pengguna: karyawan **Danamon, Adira, Home Credit** (hanya domain @danamon.co.id, @adira.co.id, @homecredit.co.id)',
    'Booking **langsung terkonfirmasi**, tanpa approval manual',
    'Aturan fair-use ditegakkan oleh sistem:',
    '  Maks. 2 booking aktif; 1 booking per fasilitas per hari',
    '  Batal paling lambat 30 menit sebelum mulai',
    '  Wajib check-in; jika tidak, otomatis no-show',
    'Admin: blokir jadwal, kelola user, domain email, dan pengaturan, semuanya tercatat di audit log',
], 16)
rect(s, Inches(8.2), Inches(1.8), Inches(4.5), Inches(4.9), GREY_LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, Inches(8.45), Inches(1.95), Inches(4), Inches(0.4), '6 fasilitas · 8 unit', 15, RED, True)
facilities = [('Tennis Court', '30 · 60 · 90 min'), ('Basketball / Futsal Court', '30 · 60 · 90 min'),
              ('Football Field', '30 · 60 · 90 min'), ('Table Tennis', '30 min'),
              ('Air Hockey (2 unit)', '30 min'), ('Foosball (2 unit)', '30 min')]
for i, (n, d) in enumerate(facilities):
    y = Inches(2.5 + i * 0.68)
    rect(s, Inches(8.45), y, Inches(4.0), Inches(0.58), WHITE, LINE, MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, Inches(8.6), y, Inches(2.6), Inches(0.58), n, 13, INK, True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(11.0), y, Inches(1.4), Inches(0.58), d, 12, GREY, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)

# 4 — Architecture -----------------------------------------------------------------------------------------
s = base('Arsitektur Sistem', 4, 'Gambaran teknis')


def box(x, y, w, h, title, body, fill=WHITE, border=LINE, tcolor=INK):
    rect(s, x, y, w, h, fill, border, MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x + Inches(0.15), y + Inches(0.1), w - Inches(0.3), Inches(0.45), title, 15, tcolor, True)
    text(s, x + Inches(0.15), y + Inches(0.55), w - Inches(0.3), h - Inches(0.6), body, 12, GREY)


def arrow(x, y, w, label, shape=MSO_SHAPE.RIGHT_ARROW, fill=RED):
    a = s.shapes.add_shape(shape, x, y, w, Inches(0.4))
    a.fill.solid()
    a.fill.fore_color.rgb = fill
    a.line.fill.background()
    text(s, x - Inches(0.2), y - Inches(0.42), w + Inches(0.4), Inches(0.4), label, 10, GREY, align=PP_ALIGN.CENTER)


# Cloudflare only serves the app files; the browser then talks to Supabase directly.
box(Inches(0.6), Inches(1.9), Inches(2.7), Inches(2.9), 'Browser / PWA',
    ['React + TypeScript', 'Dianggap TIDAK dipercaya', 'Hanya membawa kunci publik + token login user'])
box(Inches(4.3), Inches(1.7), Inches(2.6), Inches(1.45), 'Cloudflare',
    ['Static hosting saja, tanpa kode server; security headers (CSP, HSTS)'])
arrow(Inches(3.4), Inches(2.3), Inches(0.8), 'file aplikasi (HTTPS)', MSO_SHAPE.LEFT_ARROW, GREY)
arrow(Inches(3.4), Inches(3.95), Inches(4.4), 'API & Realtime langsung ke Supabase (TLS / WSS)')
rect(s, Inches(7.9), Inches(1.7), Inches(4.85), Inches(4.0), RED_SOFT, RED, MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, Inches(8.1), Inches(1.78), Inches(4.5), Inches(0.4), 'Supabase (managed) · region Singapura', 14, RED_DARK, True)
inner = [('Auth', 'Email + password, verifikasi email, hook domain'),
         ('Data API', 'Hanya fungsi schema "api"; tabel tidak terekspos'),
         ('PostgreSQL', 'Aturan bisnis, Row Level Security, audit log'),
         ('Realtime', 'Sinyal privat "jadwal berubah", tanpa data pribadi')]
for i, (t, b) in enumerate(inner):
    x = Inches(8.1 + (i % 2) * 2.33)
    y = Inches(2.3 + (i // 2) * 1.68)
    rect(s, x, y, Inches(2.2), Inches(1.5), WHITE, LINE, MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x + Inches(0.1), y + Inches(0.08), Inches(2.0), Inches(0.4), t, 13, INK, True)
    text(s, x + Inches(0.1), y + Inches(0.48), Inches(2.0), Inches(1.0), b, 11, GREY)
box(Inches(0.6), Inches(5.0), Inches(6.3), Inches(1.6), 'Email: DCU@danamon.co.id',
    ['Verifikasi & reset password via SMTP. Metode pengiriman menunggu keputusan IT Danamon (SPF/DKIM/DMARC).'])
box(Inches(7.9), Inches(5.85), Inches(4.85), Inches(0.95), 'GitHub + CI otomatis',
    ['Lint, test, dan secret scan di setiap perubahan'])

# 5 — Principles -------------------------------------------------------------------------------------------
s = base('Prinsip Keamanan Desain', 5, 'Security by design')
principles = [
    ('Browser tidak dipercaya', 'Setiap input divalidasi ulang di database seolah datang dari penyerang.'),
    ('Aturan bisnis di database', 'Tidak bisa dilewati walau API dipanggil langsung tanpa aplikasi.'),
    ('Least privilege', 'Tabel tidak terekspos; user hanya memanggil fungsi yang diizinkan.'),
    ('Tidak ada secret di browser', 'Hanya publishable key; secret key & password DB tetap di dashboard.'),
    ('Cek status secara live', 'Nonaktifkan user, ubah role, atau hapus domain berlaku seketika.'),
    ('Semua aksi admin tercatat', 'Alasan wajib dan log yang tidak bisa diubah atau dihapus.'),
]
for i, (t, b) in enumerate(principles):
    card(s, Inches(0.6 + (i % 3) * 4.1), Inches(1.85 + (i // 3) * 2.45), Inches(3.9), Inches(2.25), t, b, num=i + 1,
         body_size=14)

# 6 — Identity & access ------------------------------------------------------------------------------------
s = base('Identitas & Kontrol Akses', 6, 'Siapa boleh masuk, siapa boleh apa')
bullets(s, Inches(0.6), Inches(1.8), Inches(7.4), Inches(5), [
    'Pendaftaran hanya untuk **domain yang diizinkan**, dicek di server (Before-User-Created hook); subdomain & domain mirip ditolak',
    '**Verifikasi email wajib** sebelum bisa booking',
    'Password min. 8 karakter: huruf besar, kecil, angka, simbol; **rate limit** login & daftar',
    'Admin pertama diangkat lewat skrip terkontrol dan **tercatat di audit log**',
    'Admin tidak bisa menaikkan atau menonaktifkan dirinya sendiri',
    'Selalu ada minimal **1 admin aktif** dan **1 domain aktif**',
    'Otorisasi tidak bergantung pada klaim token; selalu dicek ke data terkini',
], 16)
rect(s, Inches(8.4), Inches(1.8), Inches(4.3), Inches(4.9), GREY_LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, Inches(8.65), Inches(1.95), Inches(3.9), Inches(0.4), 'Dua peran', 15, RED, True)
for i, (r, d) in enumerate([('USER', 'Booking, batal, check-in, lihat booking sendiri, profil'),
                            ('ADMIN', 'Semua fungsi user + blokir jadwal, kelola user, domain, pengaturan, audit log')]):
    y = Inches(2.5 + i * 2.0)
    rect(s, Inches(8.65), y, Inches(3.8), Inches(1.75), WHITE, LINE, MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, Inches(8.85), y + Inches(0.12), Inches(3.4), Inches(0.4), r, 15, INK, True)
    text(s, Inches(8.85), y + Inches(0.55), Inches(3.4), Inches(1.1), d, 13, GREY)

# 7 — Privacy ----------------------------------------------------------------------------------------------
s = base('Perlindungan Data & Privasi', 7, 'UU PDP No. 27 Tahun 2022')
rows = [
    ('Data minimal', 'Nama, email kantor, booking, notifikasi. Tidak ada nomor HP, foto, lokasi, atau device fingerprint.'),
    ('Tujuan terbatas', 'Hanya untuk booking fasilitas dan administrasi fair-use.'),
    ('Kerahasiaan', 'User hanya melihat datanya sendiri; jadwal tidak menampilkan siapa yang booking.'),
    ('Lokasi & enkripsi', 'Supabase region Singapura (disetujui PO). Enkripsi in-transit (TLS/HSTS) dan at-rest oleh provider.'),
    ('Retensi', 'Booking & audit ≥ 2 tahun (menunggu konfirmasi Legal). Notifikasi dihapus otomatis setelah 90 hari.'),
    ('Transparansi', 'Privacy notice di halaman daftar & profil; teks menunggu persetujuan Legal.'),
]
for i, (k, v) in enumerate(rows):
    y = Inches(1.8 + i * 0.82)
    rect(s, Inches(0.6), y, Inches(3.0), Inches(0.72), RED if i % 2 == 0 else RED_DARK)
    text(s, Inches(0.75), y, Inches(2.8), Inches(0.72), k, 15, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)
    rect(s, Inches(3.6), y, Inches(9.1), Inches(0.72), GREY_LIGHT)
    text(s, Inches(3.8), y, Inches(8.8), Inches(0.72), v, 14, INK, anchor=MSO_ANCHOR.MIDDLE)

# 8 — Booking integrity ------------------------------------------------------------------------------------
s = base('Integritas Booking: Anti Double-Booking', 8, 'Aturan bisnis yang tidak bisa dilewati')
bullets(s, Inches(0.6), Inches(1.8), Inches(7.6), Inches(5), [
    '**Lapis 1: penguncian baris** saat booking (profil → fasilitas → booking)',
    '**Lapis 2: constraint database.** Booking yang tumpang tindih mustahil tersimpan, bahkan lewat insert langsung',
    '**Idempotency key:** klik ganda atau jaringan putus tidak membuat booking ganda',
    'Semua waktu dihitung di server dalam **WIB**, tidak tergantung jam atau zona waktu HP',
    'Job status tiap menit (no-show, selesai) + **integrity check harian** otomatis',
], 16, gap=12)
rect(s, Inches(8.6), Inches(1.8), Inches(4.1), Inches(4.9), RED, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, Inches(8.8), Inches(2.1), Inches(3.7), Inches(0.5), 'Uji beban bersamaan', 15, RED_SOFT, True, PP_ALIGN.CENTER)
text(s, Inches(8.8), Inches(2.6), Inches(3.7), Inches(1.6), '0', 96, WHITE, True, PP_ALIGN.CENTER)
text(s, Inches(8.8), Inches(4.2), Inches(3.7), Inches(0.5), 'booking tumpang tindih', 18, WHITE, True, PP_ALIGN.CENTER)
text(s, Inches(8.8), Inches(4.9), Inches(3.7), Inches(1.6),
     ['30 user × 10 ronde', '20 request paralel × 50 iterasi', '8 / 8 skenario lulus'], 13, RED_SOFT,
     align=PP_ALIGN.CENTER)

# 9 — Web security -----------------------------------------------------------------------------------------
s = base('Keamanan Aplikasi Web', 9, 'Lapisan browser & hosting')
items = [
    ('Content Security Policy ketat', 'Tidak ada script inline atau script pihak ketiga.'),
    ('Security headers', 'HSTS, X-Frame-Options DENY, nosniff, Referrer-Policy, Permissions-Policy (kamera, mikrofon, lokasi mati).'),
    ('Anti open redirect', 'Tujuan setelah login / reset password dibatasi ke halaman internal.'),
    ('Link email aman dari scanner', 'Harus klik konfirmasi; tersedia alternatif kode 6 digit.'),
    ('Cache aman', 'Service worker hanya menyimpan file statis; data API & login tidak pernah di-cache.'),
    ('Secret scan otomatis', 'Setiap build dicek: tidak boleh ada secret / service key di bundle.'),
]
for i, (t, b) in enumerate(items):
    card(s, Inches(0.6 + (i % 2) * 6.15), Inches(1.8 + (i // 2) * 1.6), Inches(5.95), Inches(1.45), t, b, body_size=13)

# 10 — Audit & incident ------------------------------------------------------------------------------------
s = base('Audit, Monitoring & Respons Insiden', 10, 'Jejak dan kesiapan')
bullets(s, Inches(0.6), Inches(1.8), Inches(5.6), Inches(5), [
    '**Audit log append-only:** siapa, kapan, apa, sebelum/sesudah, alasan',
    'Tidak bisa diubah atau dihapus, **bahkan oleh admin**',
    'Mencakup semua tindakan admin dan siklus booking user',
    '**Monitoring:** health check job, integrity check harian, Supabase Security & Performance Advisor',
], 16, gap=12)
steps = [('Deteksi', 'Alert, integrity check, laporan user'),
         ('Kontain', 'Nonaktifkan akun, hapus domain, rotasi kunci, blokir fasilitas'),
         ('Investigasi', 'Audit log dan log Supabase'),
         ('Pulihkan', 'Restore dari backup'),
         ('Notifikasi', 'Sesuai UU PDP melalui DPO / Legal'),
         ('Review', 'Evaluasi pasca-insiden dan tindak lanjut')]
text(s, Inches(6.6), Inches(1.75), Inches(6), Inches(0.4), 'Prosedur respons insiden', 15, RED, True)
for i, (t, b) in enumerate(steps):
    y = Inches(2.25 + i * 0.75)
    c = rect(s, Inches(6.6), y, Inches(0.55), Inches(0.55), RED, shape=MSO_SHAPE.OVAL)
    p = c.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run(p, str(i + 1), 14, WHITE, True)
    text(s, Inches(7.3), y - Inches(0.02), Inches(1.7), Inches(0.6), t, 15, INK, True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(9.0), y - Inches(0.02), Inches(3.8), Inches(0.6), b, 13, GREY, anchor=MSO_ANCHOR.MIDDLE)

# 11 — Test results ----------------------------------------------------------------------------------------
s = base('Hasil Pengujian (per 14 Oktober 2026)', 11, 'Bukti, bukan klaim')
stats = [('Lulus', 'Database (pgTAP)', 'Privilege per peran, RLS, aturan bisnis, audit'),
         ('8 / 8', 'Uji konkurensi', '0 double-booking'),
         ('41 / 41', 'End-to-end', 'Android Chrome, iPhone Safari, desktop + aksesibilitas'),
         ('0 FAIL', 'OWASP ZAP baseline', '62 PASS · 5 peringatan rendah / info'),
         ('0', 'npm audit', 'Kerentanan dependency'),
         ('0 ERROR', 'Supabase Advisor', 'Project staging'),
         ('91 / 100 / 100', 'Lighthouse mobile', 'Performance / Accessibility / Best practices'),
         ('2 → 0', 'Temuan UAT', 'Temuan tampilan dari uji HP PO, diperbaiki di hari yang sama')]
for i, (big, t, b) in enumerate(stats):
    x = Inches(0.6 + (i % 4) * 3.08)
    y = Inches(1.8 + (i // 4) * 2.45)
    rect(s, x, y, Inches(2.9), Inches(2.25), GREY_LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x + Inches(0.2), y + Inches(0.15), Inches(2.6), Inches(0.8), big, 28 if len(big) < 10 else 22, RED, True)
    text(s, x + Inches(0.2), y + Inches(0.95), Inches(2.6), Inches(0.4), t, 14, INK, True)
    text(s, x + Inches(0.2), y + Inches(1.35), Inches(2.6), Inches(0.85), b, 12, GREY)

# 12 — Risk matrix -----------------------------------------------------------------------------------------
s = base('Matriks Risiko & Mitigasi', 12, 'Kemungkinan / Dampak')
risks = [
    ('Double booking akibat bug', 'Rendah', 'Tinggi', '2 lapis pengaman di database, uji konkurensi, integrity check harian'),
    ('Email verifikasi diblokir filter email perusahaan', 'Sedang', 'Tinggi', 'SMTP resmi DCU@danamon.co.id + SPF/DKIM/DMARC, allow-list di 3 perusahaan, kode 6 digit'),
    ('Pendaftaran oleh non-karyawan', 'Rendah', 'Tinggi', 'Pembatasan domain di server + verifikasi email; admin bisa nonaktifkan seketika'),
    ('Akses data user lain lewat API', 'Rendah', 'Tinggi', 'Tabel tidak terekspos, RLS, fungsi terkontrol; diuji per peran'),
    ('Penyalahgunaan hak admin', 'Rendah', 'Sedang', 'Alasan wajib, audit log permanen, rentang pengaturan dibatasi'),
    ('Kebocoran secret key', 'Rendah', 'Tinggi', 'Secret hanya di dashboard; scan bundle & repo otomatis'),
    ('Layanan berhenti / data hilang', 'Sedang', 'Tinggi', 'Paket berbayar + backup untuk produksi; status job dimonitor'),
    ('Data pribadi karyawan (UU PDP)', 'Sedang', 'Sedang', 'Data minimal, retensi, privacy notice, review Legal'),
    ('Kesalahan zona waktu di HP', 'Sedang', 'Sedang', 'Semua perhitungan WIB di server; diuji otomatis'),
]
level_color = {'Rendah': GREEN, 'Sedang': AMBER, 'Tinggi': RED}
tbl = s.shapes.add_table(len(risks) + 1, 4, Inches(0.6), Inches(1.7), Inches(12.1), Inches(5.2)).table
for c, wdt in enumerate([3.6, 1.45, 1.2, 5.85]):
    tbl.columns[c].width = Inches(wdt)
for c, h in enumerate(['Risiko', 'Kemungkinan', 'Dampak', 'Mitigasi']):
    cell = tbl.cell(0, c)
    cell.fill.solid()
    cell.fill.fore_color.rgb = INK
    cell.text_frame.paragraphs[0].text = ''
    run(cell.text_frame.paragraphs[0], h, 12, WHITE, True)
for r, (risk, lik, imp, mit) in enumerate(risks, start=1):
    for c, v in enumerate([risk, lik, imp, mit]):
        cell = tbl.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = WHITE if r % 2 else GREY_LIGHT
        cell.margin_top = cell.margin_bottom = Inches(0.04)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.text = ''
        color = level_color[v] if c in (1, 2) else (INK if c == 0 else GREY)
        run(p, v, 11, color, c in (0, 1, 2))
    tbl.rows[r].height = Inches(0.5)
tbl.rows[0].height = Inches(0.4)

# 13 — Residual risks --------------------------------------------------------------------------------------
s = base('Risiko Residual & Keterbatasan yang Diakui', 13, 'Transparansi')
residual = [
    ('Ketergantungan vendor cloud', 'Supabase & Cloudflare perlu dinilai oleh IT & Risk Danamon sesuai kebijakan bank.'),
    ('Staging memakai paket gratis', 'Produksi wajib paket berbayar dengan backup terjadwal.'),
    ('Email belum produksi', 'SMTP menunggu keputusan IT Danamon; saat ini akun uji dibuat manual oleh admin.'),
    ('Style inline di CSP', 'Dibutuhkan komponen UI; script tetap ketat. Risiko rendah.'),
    ('Data di Singapura', 'Disetujui PO; dokumentasi transfer lintas negara oleh Legal.'),
    ('Gate go-live', 'Hanya jika tidak ada temuan Critical / High yang terbuka.'),
]
for i, (t, b) in enumerate(residual):
    card(s, Inches(0.6 + (i % 3) * 4.1), Inches(1.85 + (i // 3) * 2.45), Inches(3.9), Inches(2.25), t, b,
         fill=RED_SOFT if i == 5 else GREY_LIGHT, body_size=14)

# 14 — Support needed --------------------------------------------------------------------------------------
s = base('Dukungan yang Dibutuhkan dari IT Danamon', 14, 'Permintaan')
asks = [('SMTP / relay', 'Pengirim DCU@danamon.co.id dengan SPF, DKIM, dan DMARC yang selaras.'),
        ('Allow-list email', 'Di filter email Danamon, Adira, dan Home Credit.'),
        ('Penilaian vendor & cloud', 'Supabase dan Cloudflare sesuai kebijakan bank.'),
        ('Domain resmi', 'Mis. subdomain DCU, beserta pengaturan DNS.'),
        ('Review keamanan', 'Penetration test independen sebelum go-live.'),
        ('Go-live & eskalasi', 'Kesepakatan kriteria go-live dan kontak eskalasi insiden.')]
for i, (t, b) in enumerate(asks):
    card(s, Inches(0.6 + (i % 3) * 4.1), Inches(1.85 + (i // 3) * 2.45), Inches(3.9), Inches(2.25), t, b, num=i + 1,
         body_size=14)

# 15 — Next steps ------------------------------------------------------------------------------------------
s = base('Langkah Selanjutnya', 15, 'Roadmap')
phases = [('Sekarang', 'Fase 4: UAT', 'UAT di staging oleh pengguna & admin; uji di perangkat nyata.'),
          ('Berikutnya', 'Fase 5: Produksi', 'Paket berbayar, backup, domain resmi, SMTP.'),
          ('Penutup', 'Fase 6: Serah terima', 'Panduan operasional, admin, dan prosedur insiden.')]
rect(s, Inches(0.9), Inches(2.62), Inches(11.5), Inches(0.08), LINE)
for i, (k, t, b) in enumerate(phases):
    x = Inches(0.6 + i * 4.1)
    rect(s, x + Inches(0.2), Inches(2.4), Inches(0.5), Inches(0.5), RED if i == 0 else GREY, shape=MSO_SHAPE.OVAL)
    text(s, x, Inches(1.8), Inches(3.9), Inches(0.4), k.upper(), 12, RED if i == 0 else GREY, True)
    rect(s, x, Inches(3.15), Inches(3.9), Inches(2.0), GREY_LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x + Inches(0.2), Inches(3.3), Inches(3.5), Inches(0.5), t, 17, INK, True)
    text(s, x + Inches(0.2), Inches(3.85), Inches(3.5), Inches(1.2), b, 14, GREY)
rect(s, Inches(0.6), Inches(5.5), Inches(12.1), Inches(1.1), RED, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, Inches(0.9), Inches(5.5), Inches(11.6), Inches(1.1),
     '**Gate go-live:** hasil uji lengkap · tidak ada temuan Critical/High · persetujuan IT Danamon & Product Owner',
     16, WHITE, anchor=MSO_ANCHOR.MIDDLE)

# 16 — Thank you -------------------------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
s.background.fill.solid()
s.background.fill.fore_color.rgb = INK
rect(s, 0, 0, Inches(0.25), H, RED)
text(s, Inches(1.0), Inches(2.4), Inches(11), Inches(1.2), 'Terima Kasih', 54, WHITE, True)
text(s, Inches(1.0), Inches(3.6), Inches(11), Inches(0.7), 'Diskusi & tanya jawab', 24, LINE)
text(s, Inches(1.0), Inches(5.2), Inches(11), Inches(0.5), 'Kontak: DCU@danamon.co.id', 16, LINE)

assert len(prs.slides) == TOTAL
OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print(f'Saved {OUT} ({len(prs.slides)} slides)')
